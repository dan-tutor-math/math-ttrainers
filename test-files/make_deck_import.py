# Файл для test_prompt98_deck_import.py: python3 test-files/make_deck_import.py (нужен python-pptx).
# PDF test-files/deck_import.pdf — лист самостоятельной работы, сохранён из LibreOffice.
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from PIL import Image, ImageDraw
import io

img = Image.new('RGB', (800, 500), (30, 120, 200))
d = ImageDraw.Draw(img)
for i in range(0, 800, 40): d.line([(i, 0), (i, 500)], fill=(255, 255, 255), width=3)
d.ellipse([250, 100, 550, 400], fill=(255, 200, 0))
bio = io.BytesIO(); img.save(bio, 'PNG'); bio.seek(0)

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)

# 1 title
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = 'Квадратные уравнения'
s.placeholders[1].text = 'Урок 5 · 8 класс'

# 2 bullets
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = 'План урока'
tf = s.placeholders[1].text_frame
tf.text = 'Повторим формулу дискриминанта'
p = tf.add_paragraph(); p.text = 'Решим примеры'; p.level = 1
p = tf.add_paragraph(); r = p.add_run(); r.text = 'Важно: '; r.font.bold = True; r.font.color.rgb = RGBColor(0xC0, 0, 0)
r = p.add_run(); r.text = 'D = b² − 4ac'; r.font.italic = True
p = tf.add_paragraph(); p.text = 'Домашнее задание'; p.level = 0
s.notes_slide.notes_text_frame.text = 'Спросить Петю про формулу'

# 3 shapes
s = prs.slides.add_slide(prs.slide_layouts[6])
sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3), Inches(1.5))
sh.text = 'Прямоугольник'
sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4), Inches(0.5), Inches(3), Inches(1.5))
sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(0x70, 0xAD, 0x47); sh.text = 'Скруглённый'
sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.5), Inches(0.5), Inches(2), Inches(2))
sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(0xFF, 0xC0, 0); sh.line.color.rgb = RGBColor(0, 0, 0); sh.line.width = Pt(3)
sh = s.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(10), Inches(0.5), Inches(2.5), Inches(2.5)); sh.rotation = 20
sh = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(0.5), Inches(3), Inches(3), Inches(1.2)); sh.text = 'Далее'
sh = s.shapes.add_shape(MSO_SHAPE.RECTANGULAR_CALLOUT, Inches(4.5), Inches(3), Inches(2.5), Inches(1.3)); sh.text = 'Выноска'
c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(8), Inches(3.5), Inches(11), Inches(5)); c.line.width = Pt(2.5); c.line.color.rgb = RGBColor(0xC0, 0, 0)
grp = s.shapes.add_group_shape()
g1 = grp.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(1), Inches(5), Inches(1.5), Inches(1.5))
g2 = grp.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(2.7), Inches(5), Inches(1.5), Inches(1.5)); g2.text = '∆'
tb = s.shapes.add_textbox(Inches(5), Inches(5), Inches(4), Inches(1))
tb.text_frame.text = 'Надпись без рамки'; tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)

# 4 table
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = 'Таблица значений'
t = s.shapes.add_table(4, 3, Inches(1), Inches(2), Inches(9), Inches(3)).table
for ci, h in enumerate(['x', 'y = x²', 'Заметка']): t.cell(0, ci).text = h
for ri in range(1, 4):
    t.cell(ri, 0).text = str(ri); t.cell(ri, 1).text = str(ri * ri)
t.cell(1, 2).merge(t.cell(3, 2)); t.cell(1, 2).text = 'Парабола растёт быстрее прямой'

# 5 picture + dark bg
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(0x1F, 0x2A, 0x44)
pic = s.shapes.add_picture(bio, Inches(1), Inches(1), Inches(6), Inches(3.75)); pic.rotation = -8
pic.crop_left = 0.1; pic.crop_right = 0.1
tb = s.shapes.add_textbox(Inches(7.5), Inches(1), Inches(5), Inches(2))
tb.text_frame.word_wrap = True
p = tb.text_frame.paragraphs[0]; r = p.add_run(); r.text = 'Белый текст на тёмном фоне, x'; r.font.size = Pt(32); r.font.color.rgb = RGBColor(255, 255, 255)
r = p.add_run(); r.text = '2'; r.font.size = Pt(32); r.font.color.rgb = RGBColor(255, 255, 255)
r._r.get_or_add_rPr().set('baseline', '30000')

# 6 text formatting
s = prs.slides.add_slide(prs.slide_layouts[6])
tb = s.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6), Inches(3)); tf = tb.text_frame; tf.word_wrap = True
tf.text = 'По центру и крупно'; tf.paragraphs[0].alignment = PP_ALIGN.CENTER; tf.paragraphs[0].runs[0].font.size = Pt(40)
p = tf.add_paragraph(); p.text = 'Справа, Arial'; p.alignment = PP_ALIGN.RIGHT; p.runs[0].font.name = 'Arial'
p = tf.add_paragraph(); p.text = 'Times New Roman, подчёркнутый'; p.runs[0].font.name = 'Times New Roman'; p.runs[0].font.underline = True
p = tf.add_paragraph(); p.text = 'Межстрочный 1,5 — очень длинная строка, которая переносится на несколько строк внутри надписи'; p.line_spacing = 1.5
sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7), Inches(0.5), Inches(5), Inches(3))
sh.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM; sh.text = 'Прижато к низу'
sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(0xDD, 0xEB, 0xF7); sh.line.color.rgb = RGBColor(0x2F, 0x55, 0x97)
sh.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 0, 0)
prs.save('test-files/deck_import.pptx')
print('ok')
